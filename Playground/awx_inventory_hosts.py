#!/usr/bin/env python3
"""Pobierz hosty i grupowe credentiale z AWX oraz dodaj gotowe wartości credentiali.

Wymaga Pythona 3.9+ i biblioteki openpyxl.
Instalacja zależności: python -m pip install -r requirements.txt
Grupy, ich zmienne i powiązania są pobierane zbiorczo dla każdego inventory.

Uruchom: python awx_inventory_hosts.py
Hosty z Excela: python awx_inventory_hosts.py -file input.xlsx
Diagnostyka: python awx_inventory_hosts.py -file input.xlsx -d
Plik wejściowy: kolumna host_limit z nazwami hostów z wybranego inventory.
Token: zmienna AWX_TOKEN w sekcji konfiguracji poniżej.
API: https://docs.ansible.com/projects/awx/en/latest/open_api/explorer.html
"""

import argparse
import json
import logging
import math
import re
import sys
import time
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode, urljoin, urlsplit
from urllib.request import HTTPRedirectHandler, Request, build_opener
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.exceptions import InvalidFileException
except ImportError:
    raise SystemExit(
        "Brak biblioteki openpyxl. Zainstaluj ją poleceniem: python -m pip install openpyxl"
    ) from None


# ========================== KONFIGURACJA ==========================
AWX_URL = "http://10.0.0.204:30080"
AWX_TOKEN = "XhUJQQAbseBOUzqr0sG24649GSahs6"
INVENTORY = "it2_inventory"  # Nazwa inventory albo ID jako liczba, np. 7.
# Główne zmienne grupowe pokazywane w kolumnie Credential.
CREDENTIAL_VARIABLES = ["os_credential"]
# Gotowe wartości credentiali dopisywane każdemu hostowi przed Credential.
EXTRA_CREDENTIAL_VARIABLES = ["xd","123"]  # Np. ["wartosc_1", "wartosc_2"]
FULL_CREDENTIAL_SEPARATOR = ", "  # Najpierw gotowe EXTRA, potem wartości zmiennych CREDENTIAL_VARIABLES.
OUTPUT_FILE = Path("awx_inventory_hosts.json")
EXCEL_OUTPUT_FILE = Path("awx_inventory_hosts.xlsx")
TIMEOUT = 30  # Limit czasu jednego żądania, w sekundach.
PAGE_SIZE = 200
# =================================================================

LOGGER = logging.getLogger("awx_inventory_hosts")
LOGGER.addHandler(logging.NullHandler())
LOGGER.setLevel(logging.WARNING)
LOGGER.propagate = False


class DebugFormatter(logging.Formatter):
    """Maskuj token także wtedy, gdy pojawi się w danych odpowiedzi API."""

    def __init__(self, token):
        super().__init__("[%(levelname)s] %(message)s")
        self.token = token

    def format(self, record):
        message = super().format(record)
        return message.replace(self.token, "***") if self.token else message


def configure_debug(enabled):
    handler = logging.StreamHandler(sys.stderr)
    handler.setFormatter(DebugFormatter(AWX_TOKEN.strip()))
    LOGGER.handlers.clear()
    LOGGER.addHandler(handler)
    LOGGER.setLevel(logging.DEBUG if enabled else logging.WARNING)


def debug_variables(label, values):
    if LOGGER.isEnabledFor(logging.DEBUG):
        LOGGER.debug("%s: %s", label, json.dumps(values, ensure_ascii=False, default=str))


class AWXError(RuntimeError):
    """Błąd API, konfiguracji lub danych wejściowych."""


def parse_arguments(argv=None):
    parser = argparse.ArgumentParser(
        description="Pobierz hosty i credentiale z AWX oraz zapisz wyniki w Excelu."
    )
    parser.add_argument(
        "-file", "--file", dest="input_file", type=Path, metavar="INPUT.xlsx",
        help="Użyj tylko hostów z kolumny host_limit wskazanego pliku Excel.",
    )
    parser.add_argument(
        "-d", "-debug", "--debug", action="store_true",
        help="Pokaż zmienne, cele i wyniki zapytań oraz licznik zapytań API na stderr.",
    )
    return parser.parse_args(argv)


def read_host_limits(input_file):
    """Odczytaj unikalne hosty z kolumn host_limit, zachowując kolejność z pliku."""
    LOGGER.debug("Wejście Excel: otwieram %s i szukam kolumny host_limit", input_file)
    workbook = None
    hostnames = {}
    found_column = False
    try:
        workbook = load_workbook(input_file, read_only=True, data_only=True, keep_links=False)
        for sheet in workbook.worksheets:
            LOGGER.debug("Wejście Excel: przeszukuję arkusz %r", sheet.title)
            column_index = None
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if column_index is None:
                    matching_columns = [
                        index for index, value in enumerate(row)
                        if isinstance(value, str) and value.strip().casefold() == "host_limit"
                    ]
                    if len(matching_columns) > 1:
                        raise AWXError(
                            f"Arkusz {sheet.title!r}, wiersz {row_number}: "
                            "więcej niż jedna kolumna host_limit."
                        )
                    if matching_columns:
                        column_index = matching_columns[0]
                        found_column = True
                        LOGGER.debug(
                            "Wejście Excel: znaleziono host_limit w arkuszu %r, wiersz %d, kolumna %s",
                            sheet.title, row_number, get_column_letter(column_index + 1),
                        )
                    continue

                value = row[column_index] if column_index < len(row) else None
                if value is None:
                    continue
                for hostname in re.split(r"[,;\s]+", str(value).strip()):
                    if hostname:
                        hostnames.setdefault(hostname, None)
    except (OSError, InvalidFileException, BadZipFile, ParseError) as exc:
        raise AWXError(f"Nie można odczytać pliku Excel {input_file}: {exc}") from None
    finally:
        if workbook is not None:
            workbook.close()

    if not found_column:
        raise AWXError(f"Nie znaleziono kolumny host_limit w pliku {input_file}.")
    if not hostnames:
        raise AWXError(f"Kolumna host_limit w pliku {input_file} nie zawiera żadnych hostów.")
    LOGGER.debug("Wejście Excel: odczytano %d unikalnych nazw hostów", len(hostnames))
    debug_variables(f"Hosty z pliku {input_file}", list(hostnames))
    return list(hostnames)


class NoRedirect(HTTPRedirectHandler):
    """Nie przekazuj nagłówka z tokenem przy przekierowaniu HTTP."""

    def redirect_request(self, req, fp, code, msg, headers, newurl):
        return None


class AWXClient:
    def __init__(self, base_url, token):
        self.base_url = base_url.rstrip("/") + "/"
        parsed = urlsplit(self.base_url)
        if parsed.scheme not in ("http", "https") or not parsed.netloc:
            raise AWXError("AWX_URL musi zawierać http:// albo https:// i adres serwera.")
        self.origin = (parsed.scheme, parsed.netloc)
        self.token = token
        self.opener = build_opener(NoRedirect())
        self.request_count = 0

    def get(self, path, *, purpose="Odczyt zasobu AWX"):
        url = urljoin(self.base_url, path)
        parsed = urlsplit(url)
        if (parsed.scheme, parsed.netloc) != self.origin:
            raise AWXError("API zwróciło odnośnik do innego serwera niż AWX_URL.")

        request = Request(
            url,
            headers={"Authorization": f"Bearer {self.token}", "Accept": "application/json"},
            method="GET",
        )
        self.request_count += 1
        request_number = self.request_count
        started = time.perf_counter()
        LOGGER.debug("Cel zapytania #%d: %s", request_number, purpose)
        LOGGER.debug("Zapytanie #%d: GET %s", request_number, url)
        try:
            with self.opener.open(request, timeout=TIMEOUT) as response:
                status = getattr(response, "status", None)
                payload = json.load(response)
        except HTTPError as exc:
            code = exc.code
            exc.close()
            LOGGER.debug(
                "Odpowiedź #%d: błąd HTTP %d; czas %.3f s",
                request_number, code, time.perf_counter() - started,
            )
            hint = {
                401: "Token jest nieprawidłowy lub wygasł.",
                403: "Token nie ma uprawnień do odczytu tego zasobu.",
                404: "Zasób nie istnieje albo nie jest widoczny dla tego tokena.",
                301: "Sprawdź AWX_URL: serwer przekierowuje żądanie.",
                302: "Sprawdź AWX_URL: serwer przekierowuje żądanie.",
            }.get(code, "Serwer AWX zgłosił błąd.")
            raise AWXError(f"HTTP {code}, {parsed.path}: {hint}") from None
        except (URLError, OSError) as exc:
            LOGGER.debug(
                "Odpowiedź #%d: błąd połączenia po %.3f s: %s",
                request_number, time.perf_counter() - started, exc,
            )
            raise AWXError(f"Nie można pobrać {parsed.path}: {exc}") from None
        except (ValueError, UnicodeError):
            LOGGER.debug("Odpowiedź #%d: niepoprawny JSON", request_number)
            raise AWXError(f"Odpowiedź z {parsed.path} nie jest poprawnym JSON-em.") from None

        if not isinstance(payload, dict):
            LOGGER.debug("Odpowiedź #%d: nieoczekiwany typ danych %s", request_number, type(payload).__name__)
            raise AWXError(f"Nieoczekiwany format odpowiedzi z {parsed.path}.")
        LOGGER.debug(
            "Odpowiedź #%d: HTTP %s; czas %.3f s",
            request_number, status if status is not None else "nieznany", time.perf_counter() - started,
        )
        if isinstance(payload.get("results"), list):
            LOGGER.debug(
                "Wynik zapytania #%d: rekordów na stronie=%d; łącznie według API=%s; kolejna strona=%s",
                request_number, len(payload["results"]), payload.get("count", "nie podano"),
                "tak" if payload.get("next") else "nie",
            )
        else:
            LOGGER.debug("Wynik zapytania #%d: odebrano obiekt JSON (%d pól)", request_number, len(payload))
        return payload

    def items(self, path, *, purpose="Pobieranie listy zasobów AWX", **filters):
        """Pobierz wszystkie strony, podążając za polem next z API."""
        parameters = {"page_size": PAGE_SIZE, **filters}
        debug_variables(f"Parametry: {purpose}", parameters)
        query = urlencode(parameters)
        url = urljoin(self.base_url, path) + "?" + query
        visited = set()
        page_number = 0
        record_count = 0
        while url:
            if url in visited:
                raise AWXError("API zwróciło powtarzający się odnośnik do następnej strony.")
            visited.add(url)
            page_number += 1
            page = self.get(url, purpose=f"{purpose}; strona {page_number}")
            if not isinstance(page.get("results"), list):
                raise AWXError("W odpowiedzi API brakuje listy results.")
            record_count += len(page["results"])
            yield from page["results"]
            next_page = page.get("next")
            url = urljoin(url, next_page) if next_page else None
        LOGGER.debug(
            "Zakończono: %s; pobrano %d rekordów z %d stron",
            purpose, record_count, page_number,
        )


def find_inventory(client, selector):
    if isinstance(selector, int):
        if selector < 1:
            raise AWXError("ID inventory musi być dodatnie.")
        return client.get(
            f"/api/v2/inventories/{selector}/",
            purpose=f"Pobieranie szczegółów inventory o ID {selector}",
        )

    if not isinstance(selector, str) or not selector.strip():
        raise AWXError("Ustaw INVENTORY na nazwę albo ID inventory.")
    matches = list(client.items(
        "/api/v2/inventories/", name=selector,
        purpose=f"Wyszukiwanie inventory o nazwie {selector!r}",
    ))
    if not matches:
        raise AWXError(f"Nie znaleziono dostępnego inventory o nazwie {selector!r}.")
    if len(matches) > 1:
        ids = ", ".join(str(item["id"]) for item in matches)
        raise AWXError(
            f"Nazwa {selector!r} jest niejednoznaczna. Ustaw INVENTORY na jedno z ID: {ids}."
        )
    return matches[0]


def credential_variable_names():
    """Nazwy zmiennych, których wartości są odczytywane z grup AWX."""
    variables = CREDENTIAL_VARIABLES
    if not isinstance(variables, (list, tuple)) or any(
        not isinstance(name, str) or not name.strip() for name in variables
    ):
        raise AWXError("CREDENTIAL_VARIABLES musi być listą niepustych nazw zmiennych.")
    return tuple(dict.fromkeys(name.strip() for name in variables))


def extra_credential_values():
    """Gotowe wartości z konfiguracji, wspólne dla wszystkich hostów."""
    if not isinstance(EXTRA_CREDENTIAL_VARIABLES, (list, tuple)):
        raise AWXError("EXTRA_CREDENTIAL_VARIABLES musi być listą gotowych wartości credentiali.")
    values = []
    for value in EXTRA_CREDENTIAL_VARIABLES:
        if value is None:
            continue
        if not isinstance(value, (str, int, float, bool)):
            raise AWXError("Wartości w EXTRA_CREDENTIAL_VARIABLES muszą być tekstem lub liczbą.")
        text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False)
        if text.strip() and text not in values:
            values.append(text)
    return tuple(values)


class InventoryGroupIndex:
    """Lokalny indeks grup i ich powiązań, wspólny dla wszystkich hostów."""

    def __init__(self, groups, exported, credential_variables=("os_credential",)):
        self.credential_variables = credential_variables
        self.groups = {
            group["name"]: {"id": group["id"], "name": group["name"]} for group in groups
        }
        self.direct_by_host = {}
        self.parents = {name: set() for name in self.groups}
        self.variables = {}
        self.all_groups_cache = {}

        for name in self.groups:
            group_data = exported.get(name, {})
            if not isinstance(group_data, dict) or not isinstance(group_data.get("vars", {}), dict):
                raise AWXError(f"Niepoprawne dane grupy {name!r} w eksporcie inventory.")
            # Eksport zwraca słowniki także dla zmiennych zapisanych w AWX jako YAML.
            self.variables[name] = group_data.get("vars", {})
            debug_variables(f"Zmienne grupy {name!r} (ID {self.groups[name]['id']})", self.variables[name])
            for hostname in group_data.get("hosts", []):
                self.direct_by_host.setdefault(hostname, set()).add(name)
            for child_name in group_data.get("children", []):
                if child_name not in self.parents:
                    raise AWXError(
                        f"Grupa {child_name!r} nie występuje na liście grup. "
                        "Inventory mogło się zmienić podczas odczytu; uruchom skrypt ponownie."
                    )
                self.parents[child_name].add(name)

    def all_group_names(self, name):
        """Wyznacz grupę i jej przodków raz, także przy wielu ścieżkach dziedziczenia."""
        if name not in self.all_groups_cache:
            found = set()
            pending = [name]
            while pending:
                current = pending.pop()
                if current in found:
                    continue
                found.add(current)
                pending.extend(self.parents[current] - found)
            self.all_groups_cache[name] = found
        return self.all_groups_cache[name]

    def for_host(self, hostname):
        direct_names = self.direct_by_host.get(hostname, set())
        all_names = set()
        for name in direct_names:
            all_names.update(self.all_group_names(name))
        LOGGER.debug(
            "Host %s: szukam zmiennych %s lokalnie w %d grupach (bezpośrednich=%d, nadrzędnych=%d)",
            hostname, self.credential_variables, len(all_names), len(direct_names),
            len(all_names - direct_names),
        )

        matches = []
        sorted_names = sorted(all_names)
        for variable in self.credential_variables:
            previous_matches = len(matches)
            for name in sorted_names:
                if variable in self.variables[name]:
                    LOGGER.debug(
                        "Host %s: znaleziono %s w grupie %r (%s)",
                        hostname, variable, name,
                        "bezpośrednia" if name in direct_names else "nadrzędna",
                    )
                    matches.append(
                        {
                            **self.groups[name],
                            "variable": variable,
                            "value": self.variables[name][variable],
                            "inherited": name not in direct_names,
                        }
                    )
            if len(matches) == previous_matches:
                LOGGER.debug("Host %s: nie znaleziono zmiennej %s w jego grupach", hostname, variable)
        result = {
            "groups": [self.groups[name] for name in sorted(direct_names)],
            "all_groups": [self.groups[name] for name in sorted_names],
            "os_credential_groups": [
                {
                    "id": match["id"],
                    "name": match["name"],
                    "os_credential": match["value"],
                    "inherited": match["inherited"],
                }
                for match in matches if match["variable"] == "os_credential"
            ],
            "credential_groups": matches,
        }
        debug_variables(f"Host {hostname!r}", result)
        return result


def fetch_report(client, selector, hostnames=None):
    credential_variables = credential_variable_names()
    extra_credentials = extra_credential_values()
    debug_variables("Gotowe dodatkowe wartości credentiali z konfiguracji", extra_credentials)
    inventory = find_inventory(client, selector)
    LOGGER.debug(
        "Wybrane inventory: ID=%s; nazwa=%r; typ=%s",
        inventory["id"], inventory["name"], inventory.get("kind") or "standardowe",
    )
    debug_variables("Wyszukiwane zmienne credentiali", credential_variables)
    requested = list(dict.fromkeys(hostnames)) if hostnames is not None else None
    requested_names = set(requested) if requested is not None else None
    host_filters = {"order_by": "id"}
    if requested:
        # Krótką listę filtruje AWX. Długą dopasujemy lokalnie po zbiorczym
        # odczycie, aby nie przekraczać długości URL i nie pytać o każdy host.
        candidate = {**host_filters, "page_size": PAGE_SIZE, "name__in": ",".join(requested)}
        if len(urlencode(candidate)) <= 6000:
            host_filters["name__in"] = candidate["name__in"]
            LOGGER.debug("Wyszukiwanie hostów: przekazuję %d nazw jako filtr name__in do AWX", len(requested))
        else:
            LOGGER.debug("Wyszukiwanie hostów: dopasuję %d nazw lokalnie po odczycie zbiorczym", len(requested))
    else:
        LOGGER.debug(
            "Wyszukiwanie hostów: pobieram wszystkie hosty z inventory o ID %s (nazwa: %r)",
            inventory["id"], inventory["name"],
        )
    debug_variables("Filtr hostów AWX", host_filters)
    hosts = []
    group_indexes = {}
    host_purpose = f"Pobieranie hostów z inventory o ID {inventory['id']} (nazwa: {inventory['name']!r})"
    if requested is not None:
        host_purpose += f" dla {len(requested)} nazw z pliku wejściowego"
    for host in client.items(
        f"/api/v2/inventories/{inventory['id']}/hosts/", purpose=host_purpose, **host_filters
    ):
        if requested_names is not None and host["name"] not in requested_names:
            continue
        # Hosty Smart Inventory mogą należeć do różnych inventory źródłowych.
        source_id = host.get("inventory") or inventory["id"]
        if source_id not in group_indexes:
            LOGGER.debug("Ładowanie indeksu grup źródłowego inventory o ID %s", source_id)
            groups = list(client.items(
                f"/api/v2/inventories/{source_id}/groups/", order_by="name",
                purpose=f"Pobieranie nazw i ID grup ze źródłowego inventory o ID {source_id}",
            ))
            # all=1 uwzględnia także wyłączone hosty. Bez hostvars nie pobieramy
            # niepotrzebnych zmiennych hostów. vars w eksporcie to zmienne grup.
            exported = client.get(
                f"/api/v2/inventories/{source_id}/script/?all=1",
                purpose=(
                    f"Pobieranie przypisań hostów, hierarchii i zmiennych grup ze źródłowego inventory o ID {source_id}; "
                    "uwzględniam także wyłączone hosty"
                ),
            ) if groups else {}
            group_indexes[source_id] = InventoryGroupIndex(groups, exported, credential_variables)
            LOGGER.debug(
                "Indeks grup źródłowego inventory o ID %s gotowy: %d grup i %d hostów z przypisaniami do grup; "
                "wyszukiwanie credentiali odbywa się lokalnie",
                source_id, len(groups), len(group_indexes[source_id].direct_by_host),
            )
        else:
            LOGGER.debug("Host %s: użycie indeksu grup źródłowego inventory o ID %s z pamięci", host["name"], source_id)
        hosts.append(
            {
                "id": host["id"],
                "name": host["name"],
                "enabled": host.get("enabled", True),
                **group_indexes[source_id].for_host(host["name"]),
            }
        )
    report = {
        "inventory": {"id": inventory["id"], "name": inventory["name"]},
        "host_count": len(hosts),
        "hosts": hosts,
    }
    if requested is not None:
        if not hosts:
            raise AWXError(
                f"Żaden z hostów podanych w pliku nie występuje w inventory {inventory['name']!r}."
            )
        order = {name: index for index, name in enumerate(requested)}
        hosts.sort(key=lambda host: order[host["name"]])
        found_names = {host["name"] for host in hosts}
        report["missing_hosts"] = [name for name in requested if name not in found_names]
    debug_variables("Podsumowanie hostów", {
        "host_count": len(hosts), "source_inventory_count": len(group_indexes),
        "source_inventory_ids": list(group_indexes),
        "missing_hosts": report.get("missing_hosts", []),
    })
    return report


def build_output_rows(report):
    """Przygotuj wspólne wiersze i uzupełnij hosty w JSON o dodatkowe i pełne credentiale."""
    primary_variables = credential_variable_names()
    extra_credentials = extra_credential_values()
    if not isinstance(FULL_CREDENTIAL_SEPARATOR, str):
        raise AWXError("FULL_CREDENTIAL_SEPARATOR musi być tekstem.")
    headers = ["hostname", "Extra_Credentials", "Credential", "Full_Credential"]
    extra_text = FULL_CREDENTIAL_SEPARATOR.join(extra_credentials) or "brak"
    rows = []
    for host in report["hosts"]:
        values_by_variable = {}
        full_values_by_variable = {}
        for group in host["credential_groups"]:
            if group["variable"] not in primary_variables:
                continue
            value = group["value"]
            credential = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False)
            values = values_by_variable.setdefault(group["variable"], [])
            if credential not in values:
                values.append(credential)
            if value is not None and credential.strip():
                full_values = full_values_by_variable.setdefault(group["variable"], [])
                if credential not in full_values:
                    full_values.append(credential)

        full_parts = list(extra_credentials)
        for variable in primary_variables:
            full_parts.extend(full_values_by_variable.get(variable, []))
        full_credential = FULL_CREDENTIAL_SEPARATOR.join(dict.fromkeys(full_parts)) or "brak"
        host["Extra_Credentials"] = list(extra_credentials)
        host["Full_Credential"] = full_credential
        LOGGER.debug("Host %s: dopisuję %d gotowych wartości credentiali z konfiguracji", host["name"], len(extra_credentials))
        debug_variables(f"Full_Credential hosta {host['name']!r}", full_credential)

        fields = [host["name"], extra_text]
        credentials = list(dict.fromkeys(
            value for variable in primary_variables for value in values_by_variable.get(variable, [])
        ))
        for credential in credentials or ["brak"]:
            rows.append([*fields, credential, full_credential])
    return headers, rows


def save_excel(headers, rows, output_file):
    """Zapisz wyniki jako tekst z filtrowaniem i zamrożonymi nagłówkami."""
    output_file = Path(output_file)
    LOGGER.debug("Zapis Excel: %d wierszy wyników, %d kolumn, plik %s", len(rows), len(headers), output_file)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Wyniki"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "B2"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    widths = []
    for column in sheet.iter_cols():
        longest = max(len(line) for cell in column for line in str(cell.value).split("\n"))
        width = min(70, max(18, longest + 3))
        widths.append(width)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width

    body_font = Font(name="Arial", size=10, color="1F2937")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    alternate_fill = PatternFill("solid", fgColor="F0F4F8")
    for row in sheet.iter_rows():
        line_count = 1
        for index, cell in enumerate(row):
            # Nazwy i wartości credentiali są tekstem, także gdy zaczynają się od '='.
            cell.data_type = "s"
            cell.number_format = "@"
            cell.font = header_font if cell.row == 1 else body_font
            cell.alignment = Alignment(
                horizontal="center" if cell.row == 1 else "left",
                vertical="center", wrap_text=True,
            )
            if cell.row == 1:
                cell.fill = header_fill
            elif cell.row % 2 == 0:
                cell.fill = alternate_fill
            line_count = max(line_count, sum(
                max(1, math.ceil(len(line) / (widths[index] - 3)))
                for line in str(cell.value).split("\n")
            ))
        sheet.row_dimensions[row[0].row].height = min(409, max(26, line_count * 16 + 6))
    sheet.auto_filter.ref = sheet.dimensions
    output_file.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(output_file)
    except PermissionError:
        raise AWXError(f"Nie można zapisać {output_file}. Zamknij plik w Excelu i spróbuj ponownie.") from None
    finally:
        workbook.close()


def main(argv=None):
    args = parse_arguments(argv)
    configure_debug(args.debug)
    debug_variables("Konfiguracja", {
        "AWX_URL": AWX_URL, "AWX_TOKEN": "***", "INVENTORY": INVENTORY,
        "CREDENTIAL_VARIABLES": CREDENTIAL_VARIABLES,
        "EXTRA_CREDENTIAL_VARIABLES": EXTRA_CREDENTIAL_VARIABLES,
        "FULL_CREDENTIAL_SEPARATOR": FULL_CREDENTIAL_SEPARATOR,
        "OUTPUT_FILE": OUTPUT_FILE, "EXCEL_OUTPUT_FILE": EXCEL_OUTPUT_FILE,
        "TIMEOUT": TIMEOUT, "PAGE_SIZE": PAGE_SIZE, "input_file": args.input_file,
    })
    client = None
    try:
        hostnames = None
        if args.input_file is not None:
            if args.input_file.resolve() == EXCEL_OUTPUT_FILE.resolve():
                raise AWXError(
                    "Plik wejściowy i wynikowy Excel muszą mieć różne ścieżki. "
                    "Zmień EXCEL_OUTPUT_FILE w konfiguracji."
                )
            hostnames = read_host_limits(args.input_file)

        token = AWX_TOKEN.strip()
        if not token:
            raise AWXError("Ustaw token w zmiennej AWX_TOKEN w pliku.")

        client = AWXClient(AWX_URL, token)
        report = fetch_report(client, INVENTORY, hostnames=hostnames)
        missing = report.get("missing_hosts", [])
        if missing:
            details = ", ".join(missing[:20])
            if len(missing) > 20:
                details += f" ... (+{len(missing) - 20})"
            print(f"Nie znaleziono w inventory {len(missing)} hostów: {details}", file=sys.stderr)
        headers, rows = build_output_rows(report)
        for row in rows:
            print(" , ".join(f"{header} : {value}" for header, value in zip(headers, row)))
        sys.stdout.flush()

        LOGGER.debug("Zapis JSON: %d hostów, plik %s", report["host_count"], OUTPUT_FILE)
        OUTPUT_FILE.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        save_excel(headers, rows, EXCEL_OUTPUT_FILE)
        debug_variables("Zapisane wyniki", {
            "json": OUTPUT_FILE.resolve(), "excel": EXCEL_OUTPUT_FILE.resolve(), "rows": len(rows),
        })
        return 0
    finally:
        LOGGER.debug("Łączna liczba zapytań API: %d", client.request_count if client is not None else 0)


if __name__ == "__main__":
    try:
        sys.exit(main())
    except (AWXError, OSError) as exc:
        print(f"Błąd: {exc}", file=sys.stderr)
        sys.exit(1)
    except KeyboardInterrupt:
        print("\nPrzerwano.", file=sys.stderr)
        sys.exit(130)
